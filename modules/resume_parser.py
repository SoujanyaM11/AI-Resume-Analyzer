"""Extract plain text from uploaded resumes (PDF via PyMuPDF, DOCX via python-docx).

Uploaded bytes are only ever parsed as documents; nothing is executed. Errors are raised as
``ResumeParseError`` with a user-friendly message so the UI can show them directly.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Union

import config


class ResumeParseError(Exception):
    """Raised for unsupported, oversized, unreadable or empty resumes."""


_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK"  # DOCX files are zip containers


def _clean(text: str) -> str:
    text = text.replace("\x00", " ").replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def validate_upload(filename: str, data: bytes) -> str:
    """Validate name, size and magic bytes. Returns the lowercase extension."""
    ext = Path(filename or "").suffix.lower()
    if ext not in config.ALLOWED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported file type '{ext or 'unknown'}'. Upload a PDF or DOCX resume."
        )
    if not data:
        raise ResumeParseError("The uploaded file is empty.")
    if len(data) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise ResumeParseError(f"File is larger than {config.MAX_UPLOAD_MB} MB. Upload a smaller resume.")
    magic = _PDF_MAGIC if ext == ".pdf" else _ZIP_MAGIC
    if not data.startswith(magic):
        raise ResumeParseError(f"The file does not look like a valid {ext[1:].upper()} document.")
    return ext


def _parse_pdf(data: bytes) -> str:
    try:
        import pymupdf
    except ImportError as exc:  # pragma: no cover
        raise ResumeParseError("PDF support is not installed (pip install pymupdf).") from exc
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ResumeParseError("This PDF is password-protected. Remove the password and re-upload.")
            parts = [page.get_text("text") for page in doc]
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("Could not read this PDF. It may be corrupted.") from exc
    return "\n".join(parts)


def _parse_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ResumeParseError("DOCX support is not installed (pip install python-docx).") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeParseError("Could not read this DOCX. It may be corrupted.") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # resumes often use tables for layout
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def extract_text(filename: str, data: Union[bytes, bytearray]) -> str:
    """Return cleaned resume text, or raise ``ResumeParseError``."""
    data = bytes(data)
    ext = validate_upload(filename, data)
    raw = _parse_pdf(data) if ext == ".pdf" else _parse_docx(data)
    text = _clean(raw)
    if len(text) < config.MIN_RESUME_CHARS:
        raise ResumeParseError(
            "Very little text could be extracted. Scanned or image-only resumes are not supported "
            "(no OCR). Upload a text-based PDF/DOCX or paste your resume text instead."
        )
    return text[: config.MAX_RESUME_CHARS]


def clean_pasted_text(text: str) -> str:
    """Validate resume text pasted by the user (alternative to uploading a file)."""
    text = _clean(text or "")
    if len(text) < config.MIN_RESUME_CHARS:
        raise ResumeParseError("The pasted text is too short to analyse.")
    return text[: config.MAX_RESUME_CHARS]
